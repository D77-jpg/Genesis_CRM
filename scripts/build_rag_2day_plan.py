# -*- coding: utf-8 -*-
"""Build a clean two-day RAG hardening plan DOCX for PDF export."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"D:\Trade\genesis\CRM\Genesis_CRM\docs\knowledge-rag-2day-plan.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "1F3A5F"
TEAL = "2A7F7F"
RED = "B03A2E"
AMBER = "B9770E"
GRAY = "5F6B7A"
LIGHT = "F4F6F8"
LIGHT_BLUE = "EAF2F8"
LIGHT_RED = "FDEDEC"
LIGHT_AMBER = "FEF5E7"
WHITE = "FFFFFF"
BORDER = "D5DBE3"

FONT = "Microsoft YaHei"


def set_run_font(run, size=10.5, bold=False, color=NAVY, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def style_paragraph(p, before=0, after=4, line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.alignment = align


def add_text(p, text, size=10.5, bold=False, color=NAVY, font=FONT):
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, font=font)
    return r


def add_para(doc, text="", size=10.5, bold=False, color=NAVY, before=0, after=4,
             line=1.25, align=WD_ALIGN_PARAGRAPH.LEFT, indent=None):
    p = doc.add_paragraph()
    style_paragraph(p, before=before, after=after, line=line, align=align)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def add_rich_para(doc, parts, size=10.5, before=0, after=4, line=1.25, indent=None):
    p = doc.add_paragraph()
    style_paragraph(p, before=before, after=after, line=line)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    for text, bold, color in parts:
        add_text(p, text, size=size, bold=bold, color=color)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_col_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Cm(width)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_table(doc, headers, rows, widths, header_fill=NAVY, font_size=9.5,
              first_col_bold=False, row_fills=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    no_split(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        shade(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        style_paragraph(p, before=0, after=0, line=1.1)
        add_text(p, h, size=font_size, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows):
        tr = table.add_row()
        no_split(tr)
        for c_idx, value in enumerate(row):
            cell = tr.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_fills and r_idx < len(row_fills) and row_fills[r_idx]:
                shade(cell, row_fills[r_idx])
            elif r_idx % 2 == 1:
                shade(cell, "FAFBFC")
            p = cell.paragraphs[0]
            style_paragraph(p, before=0, after=0, line=1.15)
            add_text(
                p,
                str(value),
                size=font_size,
                bold=(first_col_bold and c_idx == 0),
                color=NAVY if not (first_col_bold and c_idx == 0) else NAVY,
            )
    set_col_widths(table, widths)
    return table


def callout(doc, title, lines, fill=LIGHT_BLUE, accent=TEAL, title_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=accent, size="8")
    cell = table.rows[0].cells[0]
    cell.width = Cm(17.0)
    shade(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, left=220, right=220)
    p = cell.paragraphs[0]
    style_paragraph(p, before=0, after=5, line=1.2)
    add_text(p, title, size=11.5, bold=True, color=title_color)
    for line in lines:
        p2 = cell.add_paragraph()
        style_paragraph(p2, before=0, after=2, line=1.3)
        p2.paragraph_format.left_indent = Cm(0.2)
        add_text(p2, line, size=10, color=GRAY)


def add_section_title(doc, kicker, title, subtitle=None):
    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=3, line=1.0)
    add_text(p, kicker, size=9, bold=True, color=TEAL)
    p2 = doc.add_paragraph()
    style_paragraph(p2, before=0, after=4, line=1.05)
    add_text(p2, title, size=19, bold=True, color=NAVY)
    if subtitle:
        p3 = doc.add_paragraph()
        style_paragraph(p3, before=0, after=10, line=1.3)
        add_text(p3, subtitle, size=10.5, color=GRAY)


def add_checklist(doc, items, color=NAVY):
    for item in items:
        p = doc.add_paragraph()
        style_paragraph(p, before=0, after=3, line=1.25)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        add_text(p, "□  ", size=10.5, bold=True, color=TEAL)
        add_text(p, item, size=10.5, color=color)


def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "知识库 RAG 两天冲刺开发方案  ·  第 ", size=8, color="8A96A3")
    run = p.add_run()
    set_run_font(run, size=8, color="8A96A3")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    add_text(p, " 页", size=8, color="8A96A3")


def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, "内部开发方案  |  2026-10-26—10-27", size=8, color="8A96A3")


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)
section.different_first_page_header_footer = True

# Base styles
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.25
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.find(qn("w:rFonts"))
if rfonts is None:
    rfonts = OxmlElement("w:rFonts")
    rpr.append(rfonts)
rfonts.set(qn("w:ascii"), FONT)
rfonts.set(qn("w:hAnsi"), FONT)
rfonts.set(qn("w:eastAsia"), FONT)

add_header(section)
add_page_number_footer(section)

# ---------------- Cover ----------------
bar = doc.add_table(rows=1, cols=1)
bar.alignment = WD_TABLE_ALIGNMENT.CENTER
bar.autofit = False
set_table_borders(bar, color=NAVY, size="1")
cell = bar.rows[0].cells[0]
cell.width = Cm(17.0)
shade(cell, NAVY)
set_cell_margins(cell, top=30, bottom=30, left=0, right=0)
p = cell.paragraphs[0]
style_paragraph(p, before=0, after=0, line=1.0)
add_text(p, " ", size=1, color=WHITE)

add_para(doc, "内部开发方案 · v1.0", size=9, bold=True, color=TEAL, before=16, after=8)
add_para(doc, "知识库 RAG", size=30, bold=True, color=NAVY, before=0, after=0, line=1.0)
add_para(doc, "两天冲刺开发方案", size=30, bold=True, color=NAVY, before=0, after=10, line=1.0)
add_para(doc, "今天 + 明天｜目标：安全演示版，可内部试用", size=13, bold=False, color=TEAL, before=0, after=4)
add_para(doc, "原则：先堵风险，再补证据；两天不追求全部功能。", size=10.5, color=GRAY, before=0, after=14)

callout(doc, "先看结论", [
    "两天完成的是“安全补丁版”，不是全部五批。",
    "今天：堵住乱承诺、假引用、追问不核验。",
    "明天：补齐关键证据、历史资料不阻断、跑完回归。",
    "不做的：关键词索引、表格、产品过滤、自动训练，留到下一版。",
], fill=LIGHT_BLUE, accent=TEAL)

add_para(doc, "两天交付什么", size=12, bold=True, color=NAVY, before=16, after=7)

cards = doc.add_table(rows=2, cols=3)
cards.alignment = WD_TABLE_ALIGNMENT.CENTER
cards.autofit = False
set_table_borders(cards, color=BORDER, size="4")
card_data = [
    ("今天", "3 个必须修的漏洞", "先让系统不乱承诺", LIGHT_RED, RED),
    ("明天", "证据补全 + 发布包", "再让资料尽量完整", LIGHT_BLUE, TEAL),
    ("两天后", "可演示、可灰度", "不是全量上线", LIGHT_AMBER, AMBER),
]
for i, (tag, title, sub, fill, accent) in enumerate(card_data):
    c = cards.rows[0].cells[i]
    c.width = Cm(5.55)
    shade(c, fill)
    set_cell_margins(c, top=130, bottom=80, left=160, right=160)
    p = c.paragraphs[0]
    style_paragraph(p, before=0, after=3, line=1.0)
    add_text(p, tag, size=9, bold=True, color=accent)
    p2 = c.add_paragraph()
    style_paragraph(p2, before=0, after=2, line=1.1)
    add_text(p2, title, size=12, bold=True, color=NAVY)
    p3 = c.add_paragraph()
    style_paragraph(p3, before=0, after=0, line=1.2)
    add_text(p3, sub, size=9.5, color=GRAY)
    c2 = cards.rows[1].cells[i]
    c2.width = Cm(5.55)
    shade(c2, fill)
    set_cell_margins(c2, top=0, bottom=130, left=160, right=160)
    p4 = c2.paragraphs[0]
    style_paragraph(p4, before=0, after=0, line=1.0)
    add_text(p4, " ", size=1, color=WHITE)

add_para(doc, "测试邮件全程使用 mock 通道，不会发给真实客户。", size=9.5, bold=True, color=RED, before=14, after=0)

doc.add_page_break()

# ---------------- Today ----------------
add_section_title(doc, "今天 · 10 月 26 日", "今天：先堵住会乱承诺的口子",
                  "今天只做一件事：让没有依据的话，不能出现在可发送正文里。")

add_para(doc, "今天的排期", size=11, bold=True, color=NAVY, before=0, after=5)
add_table(doc, ["时间", "做什么", "谁", "完成标志"], [
    ("09:00–09:30", "冻结版本、备份、测试环境改成假邮件", "你 + B", "不能真发"),
    ("09:30–12:00", "把 3 个必须修的漏洞写成会失败的测试", "A", "测试先红"),
    ("13:00–15:00", "修“核验失败还保留原稿”", "A", "原稿不进发送区"),
    ("15:00–16:30", "修“无效引用提前放行”", "A", "所有引用查完才通过"),
    ("16:30–18:00", "合并工具证据 + 主题核验 + 发送门禁", "A + B", "追问也核验"),
    ("18:00–19:00", "高风险审核最小版（复用现有审核页）", "你 + A", "未批不能发"),
    ("19:00–21:00", "回归测试、修回归", "A + B", "今天验收通过"),
], widths=[2.4, 8.2, 2.0, 4.4], font_size=9.2)

add_para(doc, "今天必须过", size=11, bold=True, color=NAVY, before=13, after=5)
add_checklist(doc, [
    "核验失败、超时、断网时，正文没有未核实承诺。",
    "真假引用混在一起，不能算“有依据”。",
    "主题和正文都核验，主题不能只查长度和数字。",
    "追问中工具查到的证据，进入最终审核记录。",
    "高风险未批，不能发送；测试邮件真实发送 0。",
])

callout(doc, "今天的产出", [
    "一个能安全演示的版本：可能还会漏证据，但不会把没依据的承诺放出去。",
    "如果今天 3 个漏洞没过，明天不要进入证据开发，先修漏洞。",
], fill=LIGHT_RED, accent=RED, title_color=RED)

doc.add_page_break()

# ---------------- Tomorrow ----------------
add_section_title(doc, "明天 · 10 月 27 日", "明天：把证据补齐，做成可发布版本",
                  "明天只做一件事：让查到的资料尽量完整，让历史资料不再把查询搞崩。")

add_para(doc, "明天的排期", size=11, bold=True, color=NAVY, before=0, after=5)
add_table(doc, ["时间", "做什么", "谁", "完成标志"], [
    ("09:00–11:00", "命中窗口补前后窗口 + 脚注/整节兜底", "B", "价格/交期条件不丢"),
    ("11:00–12:00", "历史 129/500 条只警告，不阻断", "B", "正常查询能返回"),
    ("13:30–15:00", "审核/核验用完整证据，模型用精简版", "A + B", "引用位置不丢"),
    ("15:00–16:30", "回归 + 故障注入 + 权限测试", "QA + A", "关键场景全过"),
    ("16:30–17:30", "全链路假邮件测试 + 发布包/回滚", "B", "0 真实发送"),
    ("17:30–18:00", "最终验收、写限制说明", "你", "签字 / 不签字"),
], widths=[2.4, 8.2, 2.0, 4.4], font_size=9.2)

add_para(doc, "明天必须过", size=11, bold=True, color=NAVY, before=13, after=5)
add_checklist(doc, [
    "多条件问题：价格 + 交期、MOQ + 运费能一起回答。",
    "129/500 条历史资料下，正常查询不失败。",
    "每条引用能定位到具体段落，不是只显示标题。",
    "断网、超时、格式错误时，系统挂起，而不是乱答。",
    "代码检查 + 回归测试通过；回滚可用。",
    "真实发送 0，所有测试走 mock。",
])

callout(doc, "时间不够时，按这个顺序砍", [
    "1. 先保 3 个必须修的漏洞 + 历史资料不阻断。",
    "2. 证据补全只做“命中窗口 + 相邻窗口 + 脚注”，不做完整父块重构。",
    "3. 审核台先用现有页面，不新建完整系统。",
    "4. 不碰关键词索引、表格、产品过滤、自动训练。",
], fill=LIGHT_AMBER, accent=AMBER, title_color=AMBER)

# ---------------- Task list ----------------
add_section_title(doc, "任务清单", "两天任务清单（照这个做）",
                  "编号不用记，只看“完成标志”。没达到标志，就不算完成。")

add_table(doc, ["编号", "任务", "谁", "完成标志"], [
    ("P0-1", "核验失败不保留原稿", "A", "原稿只进待编辑，不进发送区"),
    ("P0-2", "无效引用不提前放行", "A", "所有引用查完才算有依据"),
    ("P0-3", "工具证据并入最终核验", "A", "追问/工具结果进入审核"),
    ("S-1", "主题核验 + 发送门禁", "A", "高风险未批不能发"),
    ("R-1", "高风险审核最小版", "你 + A", "现有审核页可批/拒/挂起"),
    ("E-1", "证据补全最小版", "B", "条件/脚注/引用位置不丢"),
    ("E-2", "历史资料不阻断", "B", "129/500 条只警告"),
    ("E-3", "审核快照不裁剪", "A + B", "审核用完整证据"),
    ("T-1", "回归/故障/权限/假邮件", "QA + B", "明天验收全过"),
], widths=[1.8, 6.9, 2.0, 6.3], font_size=9.2, first_col_bold=True)

spacer = doc.add_paragraph()
style_paragraph(spacer, before=0, after=0, line=1.0)
add_text(spacer, " ", size=1, color=WHITE)

callout(doc, "优先级", [
    "今天必须完成：P0-1、P0-2、P0-3、S-1；R-1 可复用现有审核页。",
    "明天必须完成：E-1、E-2、E-3、T-1。",
    "任何一项没过，不要用“演示能跑”代替“安全可用”。",
], fill=LIGHT_BLUE, accent=TEAL)

doc.add_page_break()

# ---------------- Not doing / next ----------------
add_section_title(doc, "范围说明", "两天不做什么，什么时候做",
                  "不是这些不重要，而是两天做不完；硬塞进来，会把已经修好的地方又弄坏。")

add_table(doc, ["不做", "为什么", "什么时候"], [
    ("FTS 持久索引（关键词索引）", "要改索引和压测，2 天做不完", "下周 3–5 天"),
    ("表格保真", "解析器改动大，容易引入新错", "下周 2–3 天"),
    ("产品软过滤 / 多 SKU", "改检索行为，要重跑评测", "下周 2–3 天"),
    ("自动训练", "没有足够真实数据，硬训会学坏", "稳定运行 4–6 周后"),
    ("300–500 题完整评测", "标注需要时间", "发布后 1–2 周"),
    ("GraphRAG / 换模型", "现在不是瓶颈，容易掩盖问题", "以后再说"),
], widths=[4.2, 7.6, 5.2], font_size=9.4, first_col_bold=True)

add_para(doc, "你的配合", size=11, bold=True, color=NAVY, before=14, after=5)
add_checklist(doc, [
    "今天 18:00：确认 A/B/C 风险清单，15 分钟。",
    "今天 20:30：抽查 5 条高风险，15 分钟。",
    "明天 17:30：最终验收，30 分钟。",
    "其他时间不用守着；系统会挂起，不会自己发。",
])

add_para(doc, "最终放行标准", size=11, bold=True, color=NAVY, before=14, after=5)
add_checklist(doc, [
    "3 个必须修的漏洞全过。",
    "高风险未批不能发。",
    "历史资料不阻断。",
    "引用能回原文。",
    "测试邮件 0 真实发送。",
    "有回滚、有备份、有说明书。",
])

callout(doc, "一句话", [
    "两天做完的不是“全部优化”，是“先不让系统说错话”。这句话最重要。",
], fill=LIGHT_RED, accent=RED, title_color=RED)

doc.save(OUT)
print(f"saved: {OUT}")
